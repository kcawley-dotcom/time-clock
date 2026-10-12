import sqlite3
import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook

DB_NAME = "work_log.db"
EXCEL_FILE = "work_log.xlsx"

def setup_db():
    conn = sqlite3.connect(DB_NAME)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS shifts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_name TEXT,
            clock_in TEXT,
            clock_out TEXT,
            total_hours TEXT
        )
    """)
    conn.commit()
    conn.close()

def current_time():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def save_shift(name, clock_in, clock_out, total_hours):
    conn = sqlite3.connect(DB_NAME)
    conn.execute(
        "INSERT INTO shifts (employee_name, clock_in, clock_out, total_hours) VALUES (?, ?, ?, ?)",
        (name, clock_in, clock_out, total_hours)
    )
    conn.commit()
    conn.close()
    excell_path = Path(EXCEL_FILE)

    if excell_path.exists():
        workbook = load_workbook(excell_path)
        if "Shifts" in workbook.sheetnames:
            sheet = workbook["Shifts"]
        else:
            sheet = workbook.create_sheet("Shifts")
            sheet.append(["Employee Name", "Clock In", "Clock Out", "Total Hours"])
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Shifts"
        sheet.append(["Employee Name", "Clock In", "Clock Out", "Hours"])

    sheet.append([name, clock_in, clock_out, total_hours])
    workbook.save(excell_path) 

def clock_in():
    name = employee_name_var.get().strip()
    if not name:
        messagebox.showwarning("Missing name", "Please enter your name.")
        return

    if status_var.get() == "CLOCKED_IN":
        messagebox.showinfo("Already clocked in", "You are already clocked in.")
        return

    start_time = current_time()
    status_var.set("CLOCKED_IN")
    status_label.config(text=f"Clocked in at: {start_time}")
    clock_in_var.set(start_time)
    clock_out_var.set("")

    messagebox.showinfo("Success", f"{name} clocked in at {start_time}")

def clock_out():
    name = employee_name_var.get().strip()
    if not name:
        messagebox.showwarning("Missing name", "Please enter your name.")
        return

    if status_var.get() != "CLOCKED_IN":
        messagebox.showinfo("Not clocked in", "You must clock in before clocking out.")
        return

    clock_in_time = datetime.strptime(clock_in_var.get(), "%Y-%m-%d %H:%M:%S")
    clock_out_time = datetime.now()

    total_seconds = (clock_out_time - clock_in_time).total_seconds()
    hours = total_seconds / 3600
    total_hours = f"{hours:.2f}"

    save_shift(name, clock_in_var.get(), clock_out_time.strftime("%Y-%m-%d %H:%M:%S"), total_hours)

    status_var.set("CLOCKED_OUT")
    status_label.config(text=f"Clocked out at: {clock_out_time.strftime('%Y-%m-%d %H:%M:%S')} | Total hours: {total_hours}")
    messagebox.showinfo("Success", f"{name} clocked out. Total hours: {total_hours}")

def view_log():
    log_window = tk.Toplevel(root)
    log_window.title("Work Log")
    log_window.geometry("700x400")

    tree = ttk.Treeview(log_window, columns=("Name", "Clock In", "Clock Out", "Hours"), show="headings")
    tree.heading("Name", text="Name")
    tree.heading("Clock In", text="Clock In")
    tree.heading("Clock Out", text="Clock Out")
    tree.heading("Hours", text="Hours")

    conn = sqlite3.connect(DB_NAME)
    rows = conn.execute("SELECT employee_name, clock_in, clock_out, total_hours FROM shifts ORDER BY id DESC").fetchall()
    conn.close()

    for row in rows:
        tree.insert("", "end", values=row)

    tree.pack(fill="both", expand=True, padx=10, pady=10)

root = tk.Tk()
root.title("Work Clock")
root.geometry("420x300")

setup_db()

employee_name_var = tk.StringVar()
clock_in_var = tk.StringVar()
clock_out_var = tk.StringVar()
status_var = tk.StringVar(value="CLOCKED_OUT")

tk.Label(root, text="Employee Name:").pack(pady=(10, 0))
tk.Entry(root, textvariable=employee_name_var, width=30).pack(pady=5)

tk.Button(root, text="Clock In", command=clock_in, width=20, bg="#4CAF50", fg="white").pack(pady=5)
tk.Button(root, text="Clock Out", command=clock_out, width=20, bg="#f44336", fg="white").pack(pady=5)
tk.Button(root, text="View Log", command=view_log, width=20).pack(pady=5)

status_label = tk.Label(root, text="Status: Clocked out", font=("Arial", 11, "bold"))
status_label.pack(pady=10)

root.mainloop()
